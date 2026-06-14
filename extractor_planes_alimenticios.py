"""
Extractor de Planes Alimenticios
Convierte archivos .docx de planes alimenticios a formato JSON estructurado.
"""

import json
import os
import re
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext, ttk
from pathlib import Path
import sys

# python-docx es una dependencia declarada en requirements.txt.
# Instálala con:  pip install -r requirements.txt
try:
    from docx import Document
except ImportError:
    sys.exit(
        "Falta la dependencia 'python-docx'.\n"
        "Instálala con:  pip install -r requirements.txt"
    )


# ── Utilidades de parseo ────────────────────────────────────────────────────

TIEMPO_MAP = {
    "desayuno":    "desayuno",
    "colación 1":  "colacion",
    "colación1":   "colacion",
    "colacion 1":  "colacion",
    "colacion1":   "colacion",
    "colación 2":  "colacion",
    "colación2":   "colacion",
    "colacion 2":  "colacion",
    "colacion2":   "colacion",
    "colación":    "colacion",
    "colacion":    "colacion",
    "comida":      "comida",
    "cena":        "cena",
}

DIAS_SEMANA = ["lunes", "martes", "miércoles", "miercoles",
               "jueves", "viernes", "sábado", "sabado", "domingo"]


def generar_id(nombre: str) -> str:
    n = nombre.lower()
    for a, b in [("á","a"),("é","e"),("í","i"),("ó","o"),("ú","u"),
                  ("ñ","n"),("à","a"),("è","e"),("ì","i"),("ò","o"),("ù","u")]:
        n = n.replace(a, b)
    n = re.sub(r"[^a-z0-9\s]", " ", n)
    n = re.sub(r"\s+", "_", n.strip())
    return n[:60]


def tiene_video(texto: str) -> bool:
    return bool(re.search(r"\bvideo\b", texto, re.IGNORECASE))


def _parsear_cantidad(s: str):
    s = s.strip().replace(",", ".")
    try:
        if "/" in s:
            partes = s.split("/")
            return round(float(partes[0]) / float(partes[1]), 4)
        v = float(s)
        return int(v) if v == int(v) else v
    except Exception:
        return None


def parsear_ingrediente_texto(texto: str) -> dict | None:
    """
    Parsea una línea de ingrediente con formato:
      Nombre (cantidad unidad)
    """
    texto = texto.strip()
    if not texto:
        return None

    m = re.match(
        r"^(?P<nombre>.+?)\s*\((?P<cant>[0-9/,.¼½¾]+[^)]*?)\)\s*$",
        texto, re.IGNORECASE
    )
    if m:
        nombre = m.group("nombre").strip().rstrip(":")
        interior = m.group("cant").strip()
        # Separar número y unidad
        m2 = re.match(r"^([0-9/,.¼½¾]+)\s*(.*?)$", interior)
        if m2:
            cant = _parsear_cantidad(m2.group(1))
            unidad = m2.group(2).strip()
        else:
            cant = None
            unidad = interior
        return {"nombre": nombre, "cantidad": cant, "unidad": unidad}

    # Sin paréntesis: toda la línea es el nombre
    if len(texto) > 2:
        return {"nombre": texto.rstrip(":").strip(), "cantidad": None, "unidad": ""}
    return None


def es_parrafo_nombre(parrafo) -> bool:
    """
    Un párrafo es NOMBRE de platillo si:
    - Su primer run con texto tiene bold=True Y el texto no empieza con bullet (•)
    Esto distingue "Lonche de huevo" (negrita) de "• Huevo entero" (bullet negrita + texto normal)
    """
    texto = parrafo.text.strip()
    if not texto:
        return False

    # Si empieza con bullet, es ingrediente
    if texto.startswith("•") or texto.startswith("-") or texto.startswith("*"):
        return False

    # Verificar si el primer run significativo es bold
    for run in parrafo.runs:
        t = run.text.strip()
        if not t:
            continue
        if t in ("•", "-", "*"):
            return False  # bullet antes del texto → ingrediente
        return bool(run.bold)

    return False


def extraer_platillos_de_celda(celda, tiempo: str) -> list[dict]:
    """Extrae uno o varios platillos de una celda de tabla."""
    platillos = []
    nombre_actual = None
    ings_actuales = []
    es_video = False
    notas_actuales = []

    def flush():
        nonlocal nombre_actual, ings_actuales, es_video, notas_actuales
        if nombre_actual:
            nombre_limpio = re.sub(r"\s*video\s*", "", nombre_actual, flags=re.IGNORECASE).strip()
            platillos.append({
                "id": generar_id(nombre_limpio),
                "nombre": nombre_limpio,
                "tiempo": tiempo,
                "video": es_video,
                "nota": "; ".join(notas_actuales) if notas_actuales else None,
                "ingredientes": ings_actuales,
            })
        nombre_actual = None
        ings_actuales = []
        es_video = False
        notas_actuales = []

    for par in celda.paragraphs:
        texto = par.text.strip()
        if not texto:
            continue

        if es_parrafo_nombre(par):
            flush()
            nombre_actual = texto
            es_video = tiene_video(texto)
        else:
            # Es un ingrediente o nota
            # Quitar el bullet si lo tiene
            linea = re.sub(r"^[•\-\*]\s*", "", texto).strip()
            if not linea:
                continue

            if nombre_actual is None:
                # No hay nombre aún: tratamos esta línea como nombre
                nombre_actual = linea
                es_video = tiene_video(linea)
                continue

            ing = parsear_ingrediente_texto(linea)
            if ing and ing["nombre"]:
                ings_actuales.append(ing)
            elif linea:
                notas_actuales.append(linea)

    flush()
    return platillos


def resolver_tiempo(texto: str) -> str:
    t = texto.lower().strip()
    for k, v in TIEMPO_MAP.items():
        if k in t:
            return v
    return "otro"


def procesar_docx(ruta: str) -> list[dict]:
    """Lee un .docx y extrae todos los platillos únicos de sus tablas."""
    doc = Document(ruta)
    platillos = []
    vistas = set()

    for tabla in doc.tables:
        if not tabla.rows:
            continue

        # Fila 0 → encabezados (tiempos)
        encabezados = [c.text.strip().lower() for c in tabla.rows[0].cells]

        for fila in tabla.rows[1:]:
            celdas = fila.cells
            if not celdas:
                continue
            primera = celdas[0].text.strip().lower()
            if not any(dia in primera for dia in DIAS_SEMANA):
                continue

            for ci, celda in enumerate(celdas[1:], 1):
                if not celda.text.strip():
                    continue
                tiempo_raw = encabezados[ci] if ci < len(encabezados) else ""
                tiempo = resolver_tiempo(tiempo_raw)

                nuevos = extraer_platillos_de_celda(celda, tiempo)
                for p in nuevos:
                    clave = (p["nombre"].lower(), p["tiempo"])
                    if clave not in vistas:
                        vistas.add(clave)
                        platillos.append(p)

    return platillos


# ── Interfaz gráfica ────────────────────────────────────────────────────────

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Extractor de Planes Alimenticios → JSON")
        self.resizable(False, False)
        self.configure(bg="#F5F7FA")
        self._build_ui()
        self.carpeta_var.trace_add("write", self._actualizar_estado)
        self._actualizar_estado()

    def _build_ui(self):
        PAD = 16
        COLOR_ACCENT = "#2E7D32"
        COLOR_BTN    = "#388E3C"
        COLOR_BTN_H  = "#1B5E20"
        COLOR_BG     = "#F5F7FA"
        COLOR_PANEL  = "#FFFFFF"

        tk.Label(self, text="🥗  Extractor de Planes Alimenticios",
                 font=("Segoe UI", 14, "bold"), bg=COLOR_BG, fg=COLOR_ACCENT).pack(pady=(PAD, 4))
        tk.Label(self, text="Lee archivos .docx y exporta un catálogo JSON de platillos",
                 font=("Segoe UI", 10), bg=COLOR_BG, fg="#555").pack(pady=(0, PAD))

        panel = tk.Frame(self, bg=COLOR_PANEL, bd=0, relief="solid",
                         highlightbackground="#DDD", highlightthickness=1)
        panel.pack(padx=PAD, pady=0, fill="x", ipadx=PAD, ipady=PAD)

        # Carpeta de entrada
        tk.Label(panel, text="📁  Carpeta con archivos Word (.docx)",
                 font=("Segoe UI", 10, "bold"), bg=COLOR_PANEL, anchor="w").pack(
                     fill="x", padx=PAD, pady=(PAD, 2))
        fila1 = tk.Frame(panel, bg=COLOR_PANEL)
        fila1.pack(fill="x", padx=PAD, pady=(0, 8))
        self.carpeta_var = tk.StringVar()
        tk.Entry(fila1, textvariable=self.carpeta_var, font=("Segoe UI", 10),
                 width=52, state="readonly", readonlybackground="#ECEFF1",
                 relief="flat", bd=1).pack(side="left", ipady=4, padx=(0, 6))
        self._btn(fila1, "Seleccionar…", self._seleccionar_carpeta,
                  COLOR_BTN, COLOR_BTN_H).pack(side="left")

        # Archivo de salida
        tk.Label(panel, text="💾  Archivo de salida (.json)",
                 font=("Segoe UI", 10, "bold"), bg=COLOR_PANEL, anchor="w").pack(
                     fill="x", padx=PAD, pady=(4, 2))
        fila2 = tk.Frame(panel, bg=COLOR_PANEL)
        fila2.pack(fill="x", padx=PAD, pady=(0, 8))
        self.salida_var = tk.StringVar()
        tk.Entry(fila2, textvariable=self.salida_var, font=("Segoe UI", 10),
                 width=52, state="readonly", readonlybackground="#ECEFF1",
                 relief="flat", bd=1).pack(side="left", ipady=4, padx=(0, 6))
        self._btn(fila2, "Seleccionar…", self._seleccionar_salida,
                  COLOR_BTN, COLOR_BTN_H).pack(side="left")

        # Botón generar
        self.btn_generar = self._btn(panel, "⚡  Generar JSON", self._generar,
                                      COLOR_ACCENT, COLOR_BTN_H,
                                      font=("Segoe UI", 11, "bold"), padx=20, pady=6)
        self.btn_generar.pack(pady=(8, PAD))

        # Progreso
        self.progreso = ttk.Progressbar(self, mode="indeterminate", length=440)
        self.progreso.pack(padx=PAD, pady=(8, 0))

        # Log
        tk.Label(self, text="Registro de actividad",
                 font=("Segoe UI", 9, "bold"), bg=COLOR_BG, anchor="w").pack(
                     fill="x", padx=PAD, pady=(8, 2))
        self.log = scrolledtext.ScrolledText(
            self, height=12, width=66, font=("Consolas", 9),
            bg="#1E272E", fg="#A8D8A8", insertbackground="white",
            relief="flat", state="disabled"
        )
        self.log.pack(padx=PAD, pady=(0, PAD))
        self.log.tag_config("ok",    foreground="#69F0AE")
        self.log.tag_config("warn",  foreground="#FFD740")
        self.log.tag_config("error", foreground="#FF5252")
        self.log.tag_config("info",  foreground="#82B1FF")

    @staticmethod
    def _btn(parent, texto, comando, bg, bg_h, font=("Segoe UI", 10), **kw):
        b = tk.Button(parent, text=texto, command=comando,
                      bg=bg, fg="white", activebackground=bg_h,
                      activeforeground="white", relief="flat",
                      cursor="hand2", font=font, **kw)
        b.bind("<Enter>", lambda e: b.config(bg=bg_h))
        b.bind("<Leave>", lambda e: b.config(bg=bg))
        return b

    def _seleccionar_carpeta(self):
        carpeta = filedialog.askdirectory(title="Seleccionar carpeta con archivos .docx")
        if carpeta:
            self.carpeta_var.set(carpeta)
            if not self.salida_var.get():
                self.salida_var.set(str(Path(carpeta) / "catalogo_platillos.json"))

    def _seleccionar_salida(self):
        ruta = filedialog.asksaveasfilename(
            title="Guardar JSON como…",
            defaultextension=".json",
            filetypes=[("JSON", "*.json"), ("Todos", "*.*")],
            initialfile="catalogo_platillos.json",
        )
        if ruta:
            self.salida_var.set(ruta)

    def _actualizar_estado(self, *_):
        tiene = bool(self.carpeta_var.get())
        self.btn_generar.config(state="normal" if tiene else "disabled",
                                bg="#2E7D32" if tiene else "#90A4AE")

    def _log(self, msg: str, tag="info"):
        self.log.config(state="normal")
        self.log.insert("end", msg + "\n", tag)
        self.log.see("end")
        self.log.config(state="disabled")

    def _generar(self):
        carpeta = self.carpeta_var.get()
        salida  = self.salida_var.get()
        if not carpeta:
            messagebox.showwarning("Falta carpeta", "Selecciona la carpeta con los archivos Word.")
            return
        if not salida:
            self._seleccionar_salida()
            salida = self.salida_var.get()
            if not salida:
                return
        self.btn_generar.config(state="disabled", text="Procesando…")
        self.progreso.start(12)
        self._log("─" * 60)
        self._log(f"📂  Carpeta:  {carpeta}", "info")
        self._log(f"💾  Salida:   {salida}", "info")
        self._log("─" * 60)
        threading.Thread(target=self._run_extraccion,
                         args=(carpeta, salida), daemon=True).start()

    def _run_extraccion(self, carpeta, salida):
        try:
            archivos = sorted(Path(carpeta).glob("*.docx"))
            if not archivos:
                self.after(0, self._log,
                           "⚠  No se encontraron archivos .docx en la carpeta.", "warn")
                return
            self.after(0, self._log, f"📋  Archivos encontrados: {len(archivos)}", "info")

            todos = []
            vistas = set()

            for arch in archivos:
                self.after(0, self._log, f"   ▶  {arch.name}", "info")
                try:
                    platillos = procesar_docx(str(arch))
                    nuevos = 0
                    for p in platillos:
                        clave = (p["nombre"].lower(), p["tiempo"])
                        if clave not in vistas:
                            vistas.add(clave)
                            todos.append(p)
                            nuevos += 1
                    self.after(0, self._log,
                               f"      ✔  {nuevos} platillo(s) extraídos", "ok")
                except Exception as e:
                    self.after(0, self._log, f"      ✗  Error: {e}", "error")

            resultado = {"platillos": todos}
            salida_path = Path(salida)
            salida_path.parent.mkdir(parents=True, exist_ok=True)
            with open(salida_path, "w", encoding="utf-8") as f:
                json.dump(resultado, f, ensure_ascii=False, indent=2)

            self.after(0, self._log, "─" * 60)
            self.after(0, self._log,
                       f"✅  {len(todos)} platillos exportados correctamente.", "ok")
            self.after(0, self._log, f"📄  Archivo: {salida}", "ok")
            self.after(0, self._log, "─" * 60)
            self.after(0, messagebox.showinfo, "¡Listo!",
                       f"Se exportaron {len(todos)} platillos.\n\nArchivo guardado en:\n{salida}")
        except Exception as e:
            self.after(0, self._log, f"❌  Error inesperado: {e}", "error")
            self.after(0, messagebox.showerror, "Error", str(e))
        finally:
            self.after(0, self._finalizar_ui)

    def _finalizar_ui(self):
        self.progreso.stop()
        self.btn_generar.config(state="normal", text="⚡  Generar JSON", bg="#2E7D32")


if __name__ == "__main__":
    app = App()
    app.mainloop()
