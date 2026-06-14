# -*- coding: utf-8 -*-
"""Carga plantillas de dieta desde archivos Word."""

from dataclasses import dataclass, field
from typing import Dict, List

from nutridietas import config
from nutridietas.herramientas import extractor_planes_alimenticios as extractor
from nutridietas.nucleo import detector_ingredientes
from nutridietas.nucleo import utilidades as U
from nutridietas.nucleo.modelos import Ingrediente, Platillo

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None


@dataclass
class ResultadoPlantillaWord:
    celdas: Dict[str, List[dict]]
    agregados: List[Platillo] = field(default_factory=list)
    conflictos: List[dict] = field(default_factory=list)


def cargar(ruta, catalogo, restricciones=None) -> ResultadoPlantillaWord:
    """Lee una plantilla .docx, llena celdas y agrega faltantes al catálogo."""
    if Document is None:
        raise RuntimeError("Falta python-docx para leer plantillas Word.")

    doc = Document(ruta)
    tabla = _buscar_tabla_dieta(doc)
    if tabla is None:
        raise ValueError("No se encontró una tabla de dieta semanal en el Word.")

    encabezados = [_texto_celda(c).strip() for c in tabla.rows[0].cells]
    idx_por_titulo = _indices_columnas(encabezados)
    celdas = {dia: [{} for _ in config.COLUMNAS] for dia in config.DIAS}

    agregados = []
    conflictos = []
    restricciones = restricciones or []
    claves_catalogo = set()
    for p in catalogo.platillos:
        claves_catalogo.update(_claves_platillo(p.nombre, p.tiempo))

    for row in tabla.rows[1:]:
        dia = _resolver_dia(_texto_celda(row.cells[0]))
        if not dia:
            continue

        for ci, col in enumerate(config.COLUMNAS):
            celda_idx = idx_por_titulo.get(col["titulo"])
            if celda_idx is None or celda_idx >= len(row.cells):
                continue

            datos = _datos_celda(row.cells[celda_idx])
            conflicto = _conflicto_celda(datos, col["tiempo"], restricciones)
            if conflicto:
                conflictos.append({
                    "dia": dia,
                    "tiempo": col["titulo"],
                    "platillo": datos.get("titulo", ""),
                    "restriccion": conflicto.restriccion,
                    "encontrado_en": conflicto.encontrado_en,
                })
                celdas[dia][ci] = {}
                continue

            celdas[dia][ci] = datos
            for platillo in _platillos_celda(row.cells[celda_idx], col["tiempo"], datos):
                claves = _claves_platillo(platillo.nombre, platillo.tiempo)
                if claves & claves_catalogo:
                    continue
                catalogo.agregar_platillo(platillo)
                claves_catalogo.update(claves)
                agregados.append(platillo)

    if agregados:
        catalogo.guardar()

    return ResultadoPlantillaWord(celdas=celdas, agregados=agregados, conflictos=conflictos)


def _buscar_tabla_dieta(doc):
    titulos = {c["titulo"] for c in config.COLUMNAS}
    for tabla in doc.tables:
        if not tabla.rows:
            continue
        encabezados = {_normalizar_titulo(c.text) for c in tabla.rows[0].cells}
        if sum(_normalizar_titulo(t) in encabezados for t in titulos) >= 3:
            return tabla
    return None


def _indices_columnas(encabezados):
    resultado = {}
    usados = set()
    for col in config.COLUMNAS:
        titulo = _normalizar_titulo(col["titulo"])
        for idx, encabezado in enumerate(encabezados):
            if idx in usados:
                continue
            if _normalizar_titulo(encabezado) == titulo:
                resultado[col["titulo"]] = idx
                usados.add(idx)
                break
    return resultado


def _resolver_dia(texto):
    texto_norm = U.normalizar(texto)
    for dia in config.DIAS:
        if U.normalizar(dia) in texto_norm:
            return dia
    return None


def _datos_celda(cell):
    lineas = _lineas_celda(cell)
    if not lineas:
        return {}

    titulo_original = _limpiar_bullet(lineas[0])
    titulo = titulo_original
    if not titulo:
        return {}

    titulo_parseado = extractor.parsear_ingrediente_texto(titulo)
    if (len(lineas) == 1 and titulo_parseado
            and titulo_parseado.get("cantidad") is not None):
        titulo = titulo_parseado["nombre"]

    ingredientes = []
    nota = None
    idx = 1
    while idx < len(lineas):
        limpia = _limpiar_bullet(lineas[idx])
        if not limpia or limpia == "+":
            idx += 1
            continue
        if _parece_nota(limpia):
            nota = limpia if nota is None else f"{nota}; {limpia}"
            idx += 1
            continue
        siguiente = _limpiar_bullet(lineas[idx + 1]) if idx + 1 < len(lineas) else ""
        if siguiente.startswith("(") and not limpia.startswith("("):
            ingredientes.extend(_ingredientes_renderizados(f"{limpia} {siguiente}", limpia))
            idx += 2
            continue
        ingredientes.extend(_ingredientes_renderizados(limpia, titulo))
        idx += 1

    if not ingredientes:
        ingredientes.extend(_ingredientes_renderizados(titulo_original, titulo))

    return {
        "titulo": _limpiar_video(titulo),
        "ingredientes": ingredientes,
        "nota": nota,
        "video": extractor.tiene_video(titulo),
    }


def _platillos_celda(cell, tiempo, datos):
    platillos = []
    if datos.get("titulo"):
        ingredientes = [_ingrediente_desde_texto(i) for i in datos.get("ingredientes", [])]
        platillos.append(Platillo(
            id=extractor.generar_id(datos["titulo"]),
            nombre=datos["titulo"],
            tiempo=tiempo,
            ingredientes=ingredientes,
            video=datos.get("video", False),
            nota=datos.get("nota"),
        ))

    for p in extractor.extraer_platillos_de_celda(cell, tiempo):
        nombre = p.get("nombre", "")
        if nombre.strip() == "+" or "\n" in nombre or "•" in nombre:
            continue
        if datos.get("titulo") and _clave(nombre) == _clave(datos["titulo"]):
            continue
        platillos.append(_platillo_desde_dict(p, tiempo))
    return platillos


def _conflicto_celda(datos, tiempo, restricciones):
    if not datos or not datos.get("titulo") or not restricciones:
        return None
    platillo = Platillo(
        id="plantilla_tmp",
        nombre=datos.get("titulo", ""),
        tiempo=tiempo,
        ingredientes=[_ingrediente_desde_texto(i) for i in datos.get("ingredientes", [])],
        video=datos.get("video", False),
        nota=datos.get("nota"),
    )
    return detector_ingredientes.detectar_conflicto_platillo(platillo, restricciones)


def _platillo_desde_dict(datos, tiempo):
    return Platillo(
        id=datos.get("id") or extractor.generar_id(datos.get("nombre", "")),
        nombre=datos.get("nombre", ""),
        tiempo=datos.get("tiempo") or tiempo,
        ingredientes=[
            Ingrediente(
                nombre=i.get("nombre", ""),
                cantidad=i.get("cantidad"),
                unidad=i.get("unidad", ""),
            )
            for i in datos.get("ingredientes", [])
            if i.get("nombre")
        ],
        video=datos.get("video", False),
        nota=datos.get("nota"),
    )


def _lineas_celda(cell):
    lineas = []
    for parrafo in cell.paragraphs:
        texto = parrafo.text.strip()
        if texto:
            lineas.extend(_separar_sublineas(texto))
    return lineas


def _separar_sublineas(texto):
    partes = []
    for segmento in texto.split("\n"):
        segmento = segmento.strip()
        if not segmento:
            continue
        subpartes = [p.strip() for p in segmento.split("•") if p.strip()]
        if segmento.startswith("•"):
            partes.extend("• " + p for p in subpartes)
        else:
            partes.append(subpartes[0] if subpartes else segmento)
            partes.extend("• " + p for p in subpartes[1:])
    return partes


def _ingredientes_renderizados(linea, titulo):
    renderizados = []
    for parte in [p.strip() for p in linea.split("/") if p.strip()]:
        texto = _limpiar_bullet(parte)
        if not texto or texto == "+":
            continue
        if texto.startswith("("):
            texto = f"{titulo} {texto}"
        renderizados.append(_ingrediente_desde_texto(texto).render())
    return renderizados


def _ingrediente_desde_texto(texto):
    datos = extractor.parsear_ingrediente_texto(texto)
    if datos:
        return Ingrediente(
            nombre=datos.get("nombre", texto),
            cantidad=datos.get("cantidad"),
            unidad=datos.get("unidad", ""),
        )
    return Ingrediente(nombre=texto)


def _texto_celda(cell):
    return "\n".join(p.text for p in cell.paragraphs)


def _limpiar_bullet(texto):
    return texto.strip().lstrip("•-*").strip()


def _limpiar_video(texto):
    return " ".join(texto.replace("VIDEO", "").replace("Video", "").split())


def _parece_nota(texto):
    normalizado = U.normalizar(texto)
    return normalizado.startswith(("se le puede", "limon +", "limon+", "nota:"))


def _normalizar_titulo(texto):
    return U.normalizar(texto)


def _clave(texto):
    return U.normalizar(texto)


def _claves_platillo(nombre, tiempo):
    claves = {(_clave(nombre), tiempo)}
    datos = extractor.parsear_ingrediente_texto(" ".join(str(nombre).split()))
    if datos and datos.get("nombre"):
        claves.add((_clave(datos["nombre"]), tiempo))
    return claves
