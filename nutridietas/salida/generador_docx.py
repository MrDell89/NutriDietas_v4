# -*- coding: utf-8 -*-
"""
generador_docx.py
=================
Convierte un PlanSemanal en un archivo Word (.docx) con EXACTAMENTE el mismo
formato del documento original del Lic. Juan Pablo Espino:

  - Pagina carta horizontal, margenes estrechos.
  - Logo arriba a la derecha.
  - Titulo centrado en verde oscuro ("Nº Plan alimenticio: Nombre").
  - Notas en la parte superior.
  - Tabla de 6 columnas (día + 5 tiempos) con cabeceras y columna de días en
    verde (#1AA27E) y texto blanco; nombres de platillos en verde oscuro,
    ingredientes con viñetas.
  - Firma del nutriologo al pie.

Usa python-docx. Las cosas que python-docx no expone directamente (sombreado
de celda, bordes, ancho fijo) se hacen manipulando el XML con OxmlElement.
"""

import logging
import os

from docx import Document
from docx.shared import Pt, Twips, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from nutridietas import config

log = logging.getLogger(__name__)


# ----------------------- helpers de bajo nivel (XML) ----------------------- #
def _sombrear(celda, color_hex):
    """Pinta el fondo de una celda."""
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def _bordes_tabla(tabla, color=config.COLOR_BORDE, size=4):
    """Bordes negros en toda la tabla (incluyendo lineas internas)."""
    tbl = tabla._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(size))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        borders.append(e)
    tblPr.append(borders)


def _fijar_anchos(tabla, anchos):
    """Fija el ancho de cada columna (twips) en modo 'layout fijo'."""
    tabla.autofit = False
    tabla.allow_autofit = False
    # tblLayout fixed
    tblPr = tabla._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for fila in tabla.rows:
        for celda, ancho in zip(fila.cells, anchos):
            celda.width = Twips(ancho)


def _margenes_celda(celda, top=40, bottom=40, left=80, right=80):
    """Margenes internos (padding) de la celda en twips."""
    tcPr = celda._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for lado, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tcPr.append(mar)


def _run(parrafo, texto, *, size, color, bold=False, italic=False):
    """Agrega un run con fuente Century Gothic, tamaño y color dados."""
    r = parrafo.add_run(texto)
    r.font.name = config.FUENTE_PRINCIPAL
    # asegurar la fuente tambien para scripts complejos
    rpr = r._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), config.FUENTE_PRINCIPAL)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = RGBColor.from_string(color)
    return r


def _sin_espacios(parrafo, antes=0, despues=2, interlineado=1.0):
    pf = parrafo.paragraph_format
    pf.space_before = Pt(antes)
    pf.space_after = Pt(despues)
    pf.line_spacing = interlineado


# ----------------------------- API principal ------------------------------ #
def generar(plan, ruta_salida=None):
    """
    Crea el .docx de un PlanSemanal y lo guarda. Devuelve la ruta del archivo.
    """
    doc = Document()

    # --- estilo por defecto: Century Gothic ---
    estilo = doc.styles["Normal"]
    estilo.font.name = config.FUENTE_PRINCIPAL
    estilo.font.size = Pt(config.TAM_INGREDIENTE)

    # --- pagina carta horizontal con margenes estrechos ---
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Twips(config.PAGINA_ANCHO)
    sec.page_height = Twips(config.PAGINA_ALTO)
    sec.top_margin = Twips(config.MARGEN_SUP)
    sec.bottom_margin = Twips(config.MARGEN_INF)
    sec.left_margin = Twips(config.MARGEN_IZQ)
    sec.right_margin = Twips(config.MARGEN_DER)

    # --- logo arriba a la derecha (en el encabezado) ---
    if os.path.exists(config.RUTA_LOGO):
        header = sec.header
        header.is_linked_to_previous = False
        p_logo = header.paragraphs[0]
        p_logo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run_logo = p_logo.add_run()
        try:
            run_logo.add_picture(config.RUTA_LOGO, height=Emu(int(1.0 * 914400)))
        except Exception:
            log.warning("No se pudo insertar el logo: %s",
                        config.RUTA_LOGO, exc_info=True)

    # --- titulo centrado ---
    p_tit = doc.add_paragraph()
    p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _sin_espacios(p_tit, antes=2, despues=6)
    ordinal_y_etq, _, nombre = plan.titulo().partition(": ")
    _run(p_tit, ordinal_y_etq + ": ", size=config.TAM_TITULO,
         color=config.COLOR_VERDE_TITULO, bold=True)
    _run(p_tit, nombre, size=config.TAM_TITULO,
         color=config.COLOR_VERDE_TITULO, bold=False)

    # --- notas superiores ---
    for nota in plan.notas_superiores:
        p_n = doc.add_paragraph()
        _sin_espacios(p_n, antes=2, despues=4)
        _run(p_n, nota, size=config.TAM_NOTA_SUP, color=config.COLOR_TEXTO)

    if plan.notas_superiores:
        doc.add_paragraph()  # un respiro antes de la tabla

    # --- tabla ---
    n_cols = 1 + len(config.COLUMNAS)
    tabla = doc.add_table(rows=1 + len(config.DIAS), cols=n_cols)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    _bordes_tabla(tabla)

    anchos = [config.ANCHO_COL_DIA] + [c["ancho"] for c in config.COLUMNAS]

    # cabecera
    fila_cab = tabla.rows[0]
    _celda_cabecera(fila_cab.cells[0], "")
    for i, col in enumerate(config.COLUMNAS, start=1):
        _celda_cabecera(fila_cab.cells[i], col["titulo"])

    # filas de dias
    for d, dia in enumerate(config.DIAS, start=1):
        fila = tabla.rows[d]
        _celda_dia(fila.cells[0], dia)
        celdas_dia = plan.celdas.get(dia, [])
        for i in range(len(config.COLUMNAS)):
            celda_obj = celdas_dia[i] if i < len(celdas_dia) else None
            _celda_contenido(fila.cells[i + 1], celda_obj)

    _fijar_anchos(tabla, anchos)

    # --- firma al pie ---
    p_firma = doc.add_paragraph()
    p_firma.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _sin_espacios(p_firma, antes=10, despues=0)
    _run(p_firma, config.FIRMA, size=config.TAM_FIRMA,
         color=config.COLOR_VERDE_TITULO, bold=False, italic=True)

    # --- guardar ---
    if ruta_salida is None:
        os.makedirs(config.CARPETA_SALIDAS, exist_ok=True)
        ruta_salida = os.path.join(config.CARPETA_SALIDAS, plan.nombre_archivo())
    doc.save(ruta_salida)
    return ruta_salida


# --------------------------- celdas especificas ---------------------------- #
def _celda_cabecera(celda, texto):
    _sombrear(celda, config.COLOR_VERDE_ENCABEZADO)
    _margenes_celda(celda, top=30, bottom=30)
    celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = celda.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _sin_espacios(p, antes=0, despues=0)
    if texto:
        _run(p, texto, size=config.TAM_CABECERA, color=config.COLOR_BLANCO, bold=True)


def _celda_dia(celda, dia):
    _sombrear(celda, config.COLOR_VERDE_ENCABEZADO)
    _margenes_celda(celda, top=30, bottom=30)
    celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = celda.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _sin_espacios(p, antes=0, despues=0)
    _run(p, dia, size=config.TAM_DIA, color=config.COLOR_BLANCO, bold=True)


def _celda_contenido(celda, celda_obj):
    _margenes_celda(celda)
    celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p0 = celda.paragraphs[0]
    _sin_espacios(p0, antes=0, despues=0)

    if celda_obj is None:
        return  # celda vacia (p.ej. Colación 2)

    titulo, ingredientes, nota = celda_obj.lineas()

    # titulo del platillo (centrado, verde oscuro, negrita)
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if titulo:
        _run(p0, titulo, size=config.TAM_PLATILLO,
             color=config.COLOR_VERDE_TITULO, bold=True)

    # caso especial "Comida libre" -> carita feliz
    if celda_obj.texto_especial:
        p_cara = celda.add_paragraph()
        p_cara.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _sin_espacios(p_cara, antes=4, despues=0)
        _run(p_cara, ":)", size=18, color=config.COLOR_TEXTO, bold=True)
        return

    # ingredientes con viñeta
    for ing in ingredientes:
        p = celda.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _sin_espacios(p, antes=0, despues=0, interlineado=1.0)
        _run(p, "• " + ing, size=config.TAM_INGREDIENTE, color=config.COLOR_TEXTO)

    # nota en cursiva (p.ej. condimentos del bistec)
    if nota:
        p = celda.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _sin_espacios(p, antes=2, despues=0)
        _run(p, nota, size=config.TAM_INGREDIENTE,
             color=config.COLOR_VERDE_TITULO, bold=True, italic=True)
